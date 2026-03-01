import os
import logging
from pathlib import Path
from typing import List, Dict, Tuple, Any, Optional
import chromadb
from chromadb.config import Settings
from bs4 import BeautifulSoup

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

try:
    import pptx
except ImportError:
    pptx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

log = logging.getLogger(__name__)

class RAGEngine:
    """Retrieval-Augmented Generation (RAG) motoru.
    Geliştirilmiş projelerin içeriğini (md, txt, py) indeksler ve benzerlik araması sunar.
    """
    def __init__(self, workspace_dir: Path, db_dir_name: str = ".rag_db"):
        self.workspace_dir = Path(workspace_dir)
        self.db_dir = self.workspace_dir / db_dir_name
        self.db_dir.mkdir(parents=True, exist_ok=True)
        
        # ChromaDB Persistent Client
        self.client = chromadb.PersistentClient(
            path=str(self.db_dir),
            settings=Settings(anonymized_telemetry=False)
        )
        
        # Collection for documents
        self.collection = self.client.get_or_create_collection(
            name="bio_ml_agent_docs",
            metadata={"hnsw:space": "cosine"}
        )
        
        self.supported_extensions = {
            ".md", ".txt", ".py", ".csv", ".json", 
            ".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm"
        }
        
    def _chunk_text(self, text: str, chunk_size: int = 1500, overlap: int = 300) -> List[str]:
        """Metni parçalara (chunk) ayırır."""
        chunks = []
        if not text:
            return chunks
            
        start = 0
        text_len = len(text)
        
        while start < text_len:
            end = min(start + chunk_size, text_len)
            
            if end < text_len:
                # Break at a logical boundary if possible
                last_boundary = -1
                for char in ('\n', '.', '?', '!', ' ', ';'):
                    pos = text.rfind(char, start + chunk_size // 2, end)
                    if pos > last_boundary:
                        last_boundary = pos
                
                if last_boundary != -1:
                    end = last_boundary + 1
                    
            chunks.append(text[start:end].strip())
            start = end - overlap
            if start >= text_len or end >= text_len:
                break
                
        return [c for c in chunks if c]

    def _load_pdf(self, path: Path) -> List[Tuple[str, Dict]]:
        """PDF dosyasını sayfa bazlı yükler."""
        if not pypdf:
            return [("[ERROR] pypdf yüklü değil.", {})]
        
        pages = []
        try:
            reader = pypdf.PdfReader(path)
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text.strip():
                    pages.append((text, {"page": i + 1, "total_pages": len(reader.pages)}))
        except Exception as e:
            log.error(f"PDF yükleme hatası ({path}): {e}")
        return pages

    def _load_docx(self, path: Path) -> List[Tuple[str, Dict]]:
        """DOCX dosyasını yükler."""
        if not docx:
            return [("[ERROR] python-docx yüklü değil.", {})]
            
        full_text = []
        try:
            doc = docx.Document(path)
            for para in doc.paragraphs:
                full_text.append(para.text)
            return [("\n".join(full_text), {})]
        except Exception as e:
            log.error(f"DOCX yükleme hatası ({path}): {e}")
            return []

    def _load_pptx(self, path: Path) -> List[Tuple[str, Dict]]:
        """PPTX dosyasını slide bazlı yükler."""
        if not pptx:
            return [("[ERROR] python-pptx yüklü değil.", {})]
            
        slides = []
        try:
            prs = pptx.Presentation(path)
            for i, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        slide_text.append(shape.text)
                text = "\n".join(slide_text)
                if text.strip():
                    slides.append((text, {"slide": i + 1, "total_slides": len(prs.slides)}))
        except Exception as e:
            log.error(f"PPTX yükleme hatası ({path}): {e}")
        return slides

    def _load_excel(self, path: Path) -> List[Tuple[str, Dict]]:
        """XLSX dosyasını sheet bazlı yükler."""
        if not openpyxl:
            return [("[ERROR] openpyxl yüklü değil.", {})]
            
        sheets = []
        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    row_str = "\t".join([str(c) if c is not None else "" for c in row])
                    if row_str.strip():
                        rows.append(row_str)
                text = "\n".join(rows)
                if text.strip():
                    sheets.append((text, {"sheet_name": sheet_name}))
        except Exception as e:
            log.error(f"Excel yükleme hatası ({path}): {e}")
        return sheets

    def _load_html(self, path: Path) -> List[Tuple[str, Dict]]:
        """HTML dosyasını yükler."""
        try:
            content = path.read_text(encoding="utf-8", errors="ignore")
            soup = BeautifulSoup(content, 'html.parser')
            # Gereksiz etiketleri temizle
            for script in soup(["script", "style"]):
                script.decompose()
            text = soup.get_text(separator='\n')
            # Fazla boşlukları temizle
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = '\n'.join(chunk for chunk in chunks if chunk)
            return [(text, {})]
        except Exception as e:
            log.error(f"HTML yükleme hatası ({path}): {e}")
            return []

    def index_workspace(self) -> int:
        """Workspace altındaki tüm desteklenen dosyaları bulur ve indeksler.
        Eskileri temizler ve baştan indeksler.
        
        Döndürdüğü değer işlenen doküman sayısıdır."""
        
        log.info("RAG indekslemesi başlatılıyor...")
        
        # Mevcut collection'ı sil ve yeniden oluştur (tam index yenileme)
        try:
            self.client.delete_collection("bio_ml_agent_docs")
        except Exception:
            pass
            
        self.collection = self.client.get_or_create_collection(
            name="bio_ml_agent_docs",
            metadata={"hnsw:space": "cosine"}
        )
        
        docs = []
        metadatas = []
        ids = []
        
        doc_count = 0
        
        # Tüm dosyaları dolaş
        for root_str, dirs, files in os.walk(self.workspace_dir):
            root = Path(root_str)
            
            # RAG DB dizinini atla
            if self.db_dir.name in root.parts or ".git" in root.parts or "venv" in root.parts:
                continue
                
            for file in files:
                file_path = root / file
                if file_path.suffix.lower() not in self.supported_extensions:
                    continue
                    
                # Çok büyük dosyaları dışla (Örn: >= 1MB csv)
                if file_path.stat().st_size > 1024 * 1024:
                    continue
                    
                try:
                    ext = file_path.suffix.lower()
                    rel_path = str(file_path.relative_to(self.workspace_dir))
                    
                    # Dosya tipine göre içeriği yükle
                    file_contents = [] # List of (text, metadata)
                    
                    if ext in {".pdf"}:
                        file_contents = self._load_pdf(file_path)
                    elif ext in {".docx"}:
                        file_contents = self._load_docx(file_path)
                    elif ext in {".pptx"}:
                        file_contents = self._load_pptx(file_path)
                    elif ext in {".xlsx"}:
                        file_contents = self._load_excel(file_path)
                    elif ext in {".html", ".htm"}:
                        file_contents = self._load_html(file_path)
                    else:
                        # Düz metin tabanlı dosyalar (py, md, txt, json, csv)
                        text = file_path.read_text(encoding="utf-8", errors="ignore")
                        file_contents = [(text, {})]

                    for content_text, extra_meta in file_contents:
                        chunks = self._chunk_text(content_text)
                        
                        for i, chunk in enumerate(chunks):
                            if not chunk.strip():
                                continue
                            
                            chunk_meta = {
                                "source": rel_path,
                                "file_type": ext,
                                "chunk_index": i
                            }
                            chunk_meta.update(extra_meta)
                            
                            docs.append(chunk)
                            metadatas.append(chunk_meta)
                            ids.append(f"{rel_path}_{len(docs)}")
                            
                    doc_count += 1
                except Exception as e:
                    log.warning(f"RAG indeksleme hatası ({file_path}): {e}")
                    
        # ChromaDB'ye ekle (batch halinde)
        if docs:
            batch_size = 100
            for i in range(0, len(docs), batch_size):
                self.collection.add(
                    documents=docs[i:i+batch_size],
                    metadatas=metadatas[i:i+batch_size],
                    ids=ids[i:i+batch_size]
                )
                
        log.info(f"RAG indeksleme tamamlandı. {doc_count} dosya işlendi, {len(docs)} parça eklendi.")
        return doc_count

    def search(self, query: str, top_k: int = 3) -> List[Dict]:
        """Verilen sorguya en benzer metin parçalarını döndürür.
        
        Returns:
            List of dicts containing 'document', 'source', 'distance'
        """
        if not query.strip():
            return []
            
        try:
            results = self.collection.query(
                query_texts=[query],
                n_results=top_k
            )
            
            if not results['documents'] or not results['documents'][0]:
                return []
                
            formatted_results = []
            for doc, meta, dist in zip(results['documents'][0], results['metadatas'][0], results['distances'][0]):
                formatted_results.append({
                    "document": doc,
                    "source": meta["source"] if meta else "unknown",
                    "distance": dist
                })
                
            return formatted_results
        except Exception as e:
            log.error(f"RAG arama hatası: {e}")
            return []
